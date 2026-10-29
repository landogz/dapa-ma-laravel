#!/usr/bin/env python3
"""Build DAPE-MA Admin UAT test scripts (Excel), matching the Phase 1 workbook layout."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "123A60"
BLUE = "055498"
GOLD = "FBD116"
WHITE = "FFFFFF"
GRAY = "BFBFBF"
ZEBRA = "F3F6FA"
RED_BG = "F8E7E8"

HERE = Path(__file__).resolve().parent
OUT = HERE / "DAPE-MA-ADMIN-UAT-TEST-SCRIPTS.xlsx"

HEADERS = [
    "Script ID",
    "Module",
    "Ref",
    "Scenario",
    "Preconditions",
    "Test Steps",
    "Test Data",
    "Expected Result",
    "Actual Result",
    "Status",
    "Severity (if Fail)",
    "Tester",
    "Date",
    "Remarks",
]
COL_WIDTHS = [12, 22, 14, 36, 32, 52, 28, 46, 30, 12, 16, 14, 12, 26]

title_font = Font(name="Arial", size=14, bold=True, color=NAVY)
label_font = Font(name="Arial", size=11, bold=True, color=NAVY)
body_font = Font(name="Arial", size=11, color="0A0A0A")
header_font = Font(name="Arial", size=10, bold=True, color=WHITE)
small_font = Font(name="Arial", size=10, color="374151")
navy_fill = PatternFill("solid", fgColor=NAVY)
gold_fill = PatternFill("solid", fgColor=GOLD)
blue_fill = PatternFill("solid", fgColor=BLUE)
zebra_fill = PatternFill("solid", fgColor=ZEBRA)
thin = Border(
    left=Side(style="thin", color=GRAY),
    right=Side(style="thin", color=GRAY),
    top=Side(style="thin", color=GRAY),
    bottom=Side(style="thin", color=GRAY),
)
wrap_top = Alignment(wrap_text=True, vertical="top")
wrap_center = Alignment(wrap_text=True, vertical="center")


def r(id_, module, ref, scenario, pre, steps, data, expected):
    return (id_, module, ref, scenario, pre, steps, data, expected)


def style_info_sheet(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    ws.row_dimensions[1].height = 22


def write_kv(ws, row, key, value):
    c1 = ws.cell(row, 1, key)
    c1.font = label_font
    c1.alignment = wrap_top
    c2 = ws.cell(row, 2, value)
    c2.font = body_font
    c2.alignment = wrap_top
    ws.row_dimensions[row].height = 48 if "\n" in str(value) else 20


def add_script_sheet(wb, title, scripts, tab_color=NAVY):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    ws.freeze_panes = "E2"
    ws.sheet_view.zoomScale = 100

    for i, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, i, h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin
        ws.column_dimensions[get_column_letter(i)].width = COL_WIDTHS[i - 1]
    ws.row_dimensions[1].height = 30

    for idx, script in enumerate(scripts, 2):
        for col, value in enumerate(script, 1):
            cell = ws.cell(idx, col, value)
            cell.font = body_font
            cell.alignment = wrap_top
            cell.border = thin
            if idx % 2 == 0:
                cell.fill = zebra_fill
        ws.row_dimensions[idx].height = 96

    last = len(scripts) + 1
    ws.auto_filter.ref = f"A1:N{last}"

    status_dv = DataValidation(
        type="list",
        formula1='"Pass,Fail,Blocked,Not Tested"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Status",
        error="Choose Pass, Fail, Blocked, or Not Tested.",
    )
    status_dv.add(f"J2:J{last}")
    sev_dv = DataValidation(
        type="list",
        formula1='"Critical,High,Medium,Low,N/A"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Severity",
        error="Choose Critical, High, Medium, Low, or N/A.",
    )
    sev_dv.add(f"K2:K{last}")
    ws.add_data_validation(status_dv)
    ws.add_data_validation(sev_dv)

    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    yellow = PatternFill("solid", fgColor="FFEB9C")
    gray = PatternFill("solid", fgColor="D9D9D9")
    ws.conditional_formatting.add(
        f"J2:J{last}",
        FormulaRule(formula=[f'J2="Pass"'], fill=green),
    )
    ws.conditional_formatting.add(
        f"J2:J{last}",
        FormulaRule(formula=[f'J2="Fail"'], fill=red),
    )
    ws.conditional_formatting.add(
        f"J2:J{last}",
        FormulaRule(formula=[f'J2="Blocked"'], fill=yellow),
    )
    ws.conditional_formatting.add(
        f"J2:J{last}",
        FormulaRule(formula=[f'J2="Not Tested"'], fill=gray),
    )
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.oddHeader.left.text = "DAPE-MA Admin UAT"
    ws.oddFooter.center.text = "&P / &N"
    return last


# ---------------------------------------------------------------------------
# Scripts
# ---------------------------------------------------------------------------

SA = [
    r("SA-001", "Sign in", "ADM-AUTH", "Open the admin sign-in page",
      "Browser; admin URL known; not signed in",
      "1. Open /admin/login\n2. Observe the Welcome back screen",
      "N/A",
      "Sign-in page loads without error. Email Address, Password, Forgot password? and Sign In are visible. Restricted admin area notice is shown."),
    r("SA-002", "Sign in", "ADM-AUTH", "Sign in with a valid Super Admin account",
      "Super Admin account exists",
      "1. Enter Email Address\n2. Enter Password\n3. Click Sign In",
      "Provisioned Super Admin email and current password",
      "User lands on Dashboard. Left panel shows the name and Super Admin. All sidebar groups are visible (Overview through Administration)."),
    r("SA-003", "Sign in", "ADM-AUTH", "Sign in with wrong password",
      "On /admin/login",
      "1. Enter a valid email\n2. Enter an incorrect password\n3. Click Sign In",
      "Wrong password",
      "Sign-in is rejected with a clear message. No admin pages are shown. Password field is not emptied of the email."),
    r("SA-004", "Sign in", "ADM-AUTH", "Blank email or password",
      "On /admin/login",
      "1. Leave Email Address blank and click Sign In\n2. Enter email, leave Password blank, click Sign In",
      "Blank fields",
      "Submit is blocked; required-field messages appear. No session is created."),
    r("SA-005", "Forgot password", "ADM-AUTH", "Request a reset code",
      "On /admin/login; account email can receive mail",
      "1. Click Forgot password?\n2. Enter the account email\n3. Click Send reset code",
      "Registered admin email",
      "Check your email step appears. A 6-digit code is sent from DAPE-MA. Back to Login returns to the sign-in form."),
    r("SA-006", "Forgot password", "ADM-AUTH", "Reset password with a valid new password",
      "Reset code received",
      "1. Enter the 6-digit code\n2. Enter New password and Confirm that meet the policy\n3. Save\n4. Sign in with the new password",
      "Code from email; new password e.g. Sample#2026Aa",
      "Password is updated. Sign-in with the new password succeeds. The old password no longer works."),
    r("SA-007", "Forgot password", "ADM-AUTH", "Weak or mismatched new password is rejected",
      "On Reset password step",
      "1. Enter a short password (e.g. 12345678)\n2. Try password without a symbol\n3. Enter a valid password with a different confirmation",
      "Weak / mismatched passwords",
      "Each case is blocked. Hint is shown: at least 8 characters, upper and lower case, a number, a symbol. Password is not changed."),
    r("SA-008", "Profile", "ADM-AUTH", "Edit own name and photo",
      "Signed in as Super Admin",
      "1. Click the account menu (top right)\n2. Click Edit Profile\n3. Change First Name / Last Name\n4. Choose Photo (optional)\n5. Click Save Changes",
      "Updated name; JPEG/PNG under 5 MB",
      "Toast confirms save. Sidebar and header show the new name/photo. Email and Role remain read-only."),
    r("SA-009", "Profile", "ADM-AUTH", "Change own password from Edit Profile",
      "On Edit Profile",
      "1. Enter Current Password\n2. Enter New Password and Confirm meeting the policy\n3. Click Save Changes (single primary button; leave password blank to keep it)\n4. Sign Out and sign in with the new password",
      "Current password; new policy-compliant password",
      "Password updates via Save Changes. New password signs in. Reuse of the current password or last 5 passwords is rejected. There is no separate Update Password button."),
    r("SA-010", "Profile", "ADM-AUTH", "Sign out",
      "Signed in",
      "1. Open the account menu\n2. Click Sign Out\n3. Try to open /admin/dashboard",
      "N/A",
      "User returns to sign-in. Protected admin pages are not shown without signing in again."),
    r("SA-011", "Access", "ADM-AUTH", "Mobile app user cannot open admin",
      "An App User (mobile) account exists",
      "1. Open /admin/login\n2. Sign in with the app user email and password",
      "App user credentials",
      "Access is denied. No admin menus or data are shown."),
    r("SA-012", "Dashboard", "ADM-DASH", "Dashboard overview loads",
      "Signed in as Super Admin",
      "1. Open Overview → Dashboard\n2. Wait for counts to finish loading",
      "N/A",
      "Quick Actions, count cards (Posts, Rehab Centers, Contests, IEC, Translations, Notifications, Users, Journal), activity chart and recent activity display without error."),
    r("SA-013", "Dashboard", "ADM-DASH", "Quick Action opens Create Post",
      "On Dashboard",
      "1. Click New Post (Quick Actions)\n2. Close Cancel without saving",
      "N/A",
      "Create Post dialog opens (Super Admin: Create & Publish). Cancel closes it. No stray post is created."),
    r("SA-014", "Dashboard", "ADM-DASH", "Mega search finds a page",
      "On any admin page",
      "1. Click Search pages, actions, and records… (or ⌘K)\n2. Type posts\n3. Open the Posts result",
      "posts",
      "Results include the Posts page. Opening it navigates to /admin/posts."),
    r("SA-015", "Dashboard", "ADM-DASH", "Mega search finds a record",
      "At least one post exists",
      "1. Open mega search\n2. Type part of a known post title\n3. Open the matching record",
      "Known post title fragment",
      "Matching records appear after 2+ characters. Opening one goes to the related admin page/record."),
    r("SA-016", "Lists", "ADM-LIST", "Search, sort and pagination on Posts",
      "On Posts; more than 10 posts if possible",
      "1. Type in Search posts\n2. Click Title to sort\n3. Change Rows per page\n4. Use pagination if shown",
      "Search term from an existing title",
      "List filters as you search. Sort reverses on second click. Pagination and row count work. No page reload."),
    r("SA-017", "Lists", "ADM-LIST", "Actions tooltips and More actions",
      "On Posts desktop width",
      "1. Hover View and Edit icons\n2. Open More actions (…)",
      "N/A",
      "Hover shows the tooltip label (View, Edit, …). Extra commands are in More actions. Actions stay on the right."),
    r("SA-018", "Lists", "ADM-LIST", "Right-click opens the same Actions",
      "On Posts; a row exists",
      "1. Right-click a post row\n2. Choose View from the menu",
      "Any post row",
      "Browser menu is suppressed. Context menu lists the same commands as Actions. View opens the post dialog."),
    r("SA-019", "Lists", "ADM-LIST", "Phone-width layout stays usable",
      "Browser width ~375 px (or device toolbar iPhone)",
      "1. Narrow the window\n2. Open the ☰ menu\n3. Open Posts and confirm cards/buttons",
      "N/A",
      "Sidebar opens from ☰. List becomes cards (not one-letter wrapping). Primary buttons remain tappable (~44 px). No horizontal page scroll."),
    r("SA-020", "Posts", "ADM-POSTS", "Super Admin creates and publishes a post",
      "On Posts",
      "1. Click Create Post\n2. Optional Featured Image (JPG/PNG/WEBP/GIF ≤5 MB)\n3. Enter Title\n4. Write Body with a numbered list and a bullet list\n5. Choose Category\n6. Leave Comments Enabled\n7. Click Create & Publish",
      "Title: UAT Post SA-020; Body with 1. One 2. Two and a bullet",
      "Toast confirms. Post appears as Published. Numbered list shows 1, 2, 3 in the editor and when viewed. Post is eligible for the mobile app."),
    r("SA-021", "Posts", "ADM-POSTS", "Featured post flag",
      "Create/Edit Post dialog",
      "1. Tick Featured post\n2. Save/publish\n3. Confirm Featured badge on the list",
      "Existing or new post",
      "Featured badge shows on the row. Featured posts are intended for the mobile carousel when more than one is selected."),
    r("SA-022", "Posts", "ADM-POSTS", "Required Body blocks save",
      "Create Post dialog open",
      "1. Enter Title only\n2. Leave Body empty\n3. Click Create & Publish",
      "Title only",
      "Save is blocked with a validation message. No published post is created."),
    r("SA-023", "Posts", "ADM-POSTS", "View post, comments and delete a comment",
      "A published post with at least one user comment",
      "1. Click View\n2. Open Comments\n3. Click Delete on a comment\n4. Confirm",
      "Existing comment",
      "View shows body as HTML (not raw tags). Delete asks for confirmation. After confirm, comment is gone and a toast appears."),
    r("SA-024", "Posts", "ADM-POSTS", "Disable comments from Actions",
      "A post with comments enabled",
      "1. Open More actions\n2. Click Disable comments",
      "Any published post",
      "Comments badge/status becomes Disabled. Users cannot add new comments on that post in the app."),
    r("SA-025", "Posts", "ADM-POSTS", "Edit a published post",
      "A published post exists",
      "1. Click Edit\n2. Change the title\n3. Save",
      "Updated title",
      "Toast confirms. List shows the new title. App shows the update after refresh."),
    r("SA-026", "Posts", "ADM-POSTS", "Approve & Schedule a pending post",
      "A post in Pending review (create via Editor if needed)",
      "1. Open More actions → Approve & Schedule\n2. Set a go-live a few minutes ahead\n3. Confirm\n4. Wait and refresh",
      "Future date/time",
      "Status becomes Scheduled, then Published at the time. Post appears in the app after it goes live."),
    r("SA-027", "Posts", "ADM-POSTS", "Publish now from scheduled/pending",
      "Pending review or Scheduled post",
      "1. Click Publish now\n2. Confirm",
      "Pending/scheduled post",
      "Status becomes Published immediately. Toast confirms."),
    r("SA-028", "Posts", "ADM-POSTS", "Archive a published post",
      "A published post",
      "1. More actions → Archive\n2. Confirm\n3. Check the mobile app Featured/list",
      "Published post",
      "Confirmation dialog appears. After confirm, status is Archived and the post disappears from the app."),
    r("SA-029", "Posts", "ADM-POSTS", "Delete a post",
      "A post that may be deleted (use the UAT post)",
      "1. More actions → Delete\n2. Confirm",
      "UAT post",
      "Confirmation is required. Post is removed from the list. Toast confirms. Cannot be undone."),
    r("SA-030", "Contests", "ADM-CONTESTS", "Add a Song contest",
      "On Contests",
      "1. Click Add Contest\n2. Category = Song Contest\n3. Allowed types = Song + Playlist\n4. Enter title, year, start/end dates\n5. Status = Open for submissions\n6. Visibility = Active in app\n7. Add rules\n8. Create",
      "Title: UAT Song Contest",
      "Contest appears in the list as Song Contest / Open. Toast confirms. Visible in the app when Active."),
    r("SA-031", "Contests", "ADM-CONTESTS", "Filter by type and status",
      "At least one contest per type if possible",
      "1. Click Song Contest, then All types\n2. Click Open for submissions, then All statuses",
      "N/A",
      "List and counts update without a page reload. All types / All statuses restore the full list."),
    r("SA-032", "Contests", "ADM-CONTESTS", "Review submissions",
      "A contest with at least one entry (or empty state)",
      "1. Click the submissions/review action\n2. Observe the entries dialog",
      "Any contest",
      "Submissions dialog opens. Empty state is clear if none. No error."),
    r("SA-033", "Contests", "ADM-CONTESTS", "Edit and delete contest",
      "UAT contest exists",
      "1. Edit the title and save\n2. Delete the contest and confirm",
      "UAT Song Contest",
      "Edit toast then updated title. Delete warns that submissions are also deleted. After confirm, row is gone."),
    r("SA-034", "IEC Materials", "ADM-IEC", "Add IEC material",
      "On IEC Materials",
      "1. Click Add IEC Material\n2. Enter title and topic (e.g. Youth)\n3. Choose Image or YouTube\n4. Upload/paste media\n5. Save",
      "Sample image or YouTube URL",
      "Item appears with preview thumbnail/topic. Toast confirms."),
    r("SA-035", "IEC Materials", "ADM-IEC", "Preview, edit, delete IEC",
      "IEC item exists",
      "1. Click Preview\n2. Close\n3. Edit title and save\n4. Delete and confirm",
      "UAT IEC item",
      "Preview shows the media. Edit updates the list. Delete requires confirm and removes the row."),
    r("SA-036", "Rehab Centers", "ADM-REHAB", "Add rehab center with address suggestion and map pin",
      "On Rehab Centers",
      "1. Click Add Rehab Center\n2. Enter Name\n3. Type at least 3 letters in Address\n4. Select a Philippine suggestion\n5. Confirm the map pin moved\n6. Optional contact/website\n7. Status Active\n8. Create",
      "Name: UAT Rehab; Address e.g. Quezon City",
      "Suggestions appear on type. After select, dropdown stays closed until the text changes. Pin is set. Create succeeds. Pin shows on the facility map."),
    r("SA-037", "Rehab Centers", "ADM-REHAB", "Cannot save without a map pin",
      "Add Rehab Center dialog",
      "1. Enter Name and a free-typed Address without picking a suggestion or pin\n2. Click Create",
      "Name only / unpinned address",
      "Save is blocked (pin / region / province required). No listing is created."),
    r("SA-038", "Rehab Centers", "ADM-REHAB", "Edit and delete rehab center",
      "UAT Rehab exists",
      "1. Edit name and save\n2. Delete and confirm",
      "UAT Rehab",
      "Edit toast and updated name. Delete confirmation then row and map pin are gone."),
    r("SA-039", "Care Toolkit", "ADM-TOOLKIT", "Add a Stress Check question (EN/TL)",
      "On Care Toolkit",
      "1. Click Add Question\n2. Toolkit = Stress Check\n3. Answer type = 5-point scale\n4. Enter Question (English) and Question (Tagalog)\n5. Save",
      "EN and TL question text",
      "Question appears under Stress Check. Toast confirms. App Self-care Toolkit can show it after refresh."),
    r("SA-040", "Care Toolkit", "ADM-TOOLKIT", "Filter by toolkit type",
      "Questions exist in more than one toolkit",
      "1. Click Stress Check, Anxiety Check, Sleep Quality, then All toolkits",
      "N/A",
      "Counts and list match the filter. No page reload."),
    r("SA-041", "Care Toolkit", "ADM-TOOLKIT", "Edit and delete a question",
      "UAT question exists",
      "1. Edit the English text and save\n2. Delete and confirm",
      "UAT toolkit question",
      "Edit updates the row. Delete requires confirm and removes the question."),
    r("SA-042", "Care Checks", "ADM-CHECKS", "View submitted care checks",
      "At least one mobile check exists (or empty state)",
      "1. Open Care → Care Checks\n2. Switch 7 / 30 / 90 days\n3. Filter All checks / Stress / Anxiety / Sleep and Lower/Higher tier",
      "N/A",
      "Disclaimer is shown (not a clinical diagnosis). Cards and list load. Filters work. Empty state is clear if none."),
    r("SA-043", "Care Checks", "ADM-CHECKS", "Export care checks CSV",
      "On Care Checks",
      "1. Click Export CSV\n2. Open the downloaded file",
      "N/A",
      "CSV downloads. Columns include date, user, check, level/tier, score. Toast confirms."),
    r("SA-044", "Care Support", "ADM-SUPPORT", "Add a 24/7 hotline with a Philippine number",
      "On Care Support",
      "1. Click Add Resource\n2. Category = 24/7 Hotline\n3. Titles EN/TL\n4. Phone 09171234567\n5. Save",
      "09171234567",
      "Resource appears under 24/7 Hotlines. Toast confirms."),
    r("SA-045", "Care Support", "ADM-SUPPORT", "Reject a non-Philippine phone number",
      "Add/Edit Resource as Hotline",
      "1. Enter +1 2025550123 as phone\n2. Try to save",
      "+1 2025550123",
      "Save is blocked. Message requires a Philippine number. Record is not created."),
    r("SA-046", "Care Support", "ADM-SUPPORT", "Hotline short codes and multiple numbers",
      "Add Resource as Hotline",
      "1. Enter 1553|911 in the phone field (if shown with short-code support)\n2. Save",
      "1553|911",
      "Short codes and | separators are accepted on Care Support hotlines. Resource saves."),
    r("SA-047", "Care Support", "ADM-SUPPORT", "Edit and delete a support resource",
      "UAT hotline exists",
      "1. Edit the English title and save\n2. Delete and confirm",
      "UAT hotline",
      "Edit toast. Delete confirmation then row is gone."),
    r("SA-048", "Journal", "ADM-JOURNAL", "Filter journal by user and sky",
      "At least one journal entry exists",
      "1. Open Care → Journal\n2. Search writers and select one user\n3. Select a sky color\n4. Return to All users / All skies",
      "Existing writer/sky",
      "List filters correctly. Avatars/swatches display. Counts update."),
    r("SA-049", "Journal", "ADM-JOURNAL", "View a journal entry",
      "An entry exists",
      "1. Click View on a row",
      "Any entry",
      "Dialog shows date, sky, feelings and notes. No crash."),
    r("SA-050", "Journal", "ADM-JOURNAL", "Delete a journal entry",
      "A disposable UAT entry exists",
      "1. Click Delete\n2. Confirm",
      "UAT entry",
      "Confirmation dialog appears. After confirm, entry is gone and a toast shows."),
    r("SA-051", "Hope Directory", "ADM-HOPE-DIR", "Add organization with address autocomplete and PH phone",
      "On Hope Directory",
      "1. Click Add Organization\n2. Enter Name\n3. Category Gov't / NGOs / Schools\n4. Type Address and pick a PH suggestion\n5. Enter phone 09181234567\n6. Save",
      "Name: UAT Partner; 09181234567",
      "Organization appears under the category. Address suggestion works. Phone accepted. Toast confirms."),
    r("SA-052", "Hope Directory", "ADM-HOPE-DIR", "Filter Hope Directory by category",
      "Organizations in more than one category if possible",
      "1. Click Gov't, NGOs, Schools, then All categories",
      "N/A",
      "List and counts match the tab."),
    r("SA-053", "Hope Directory", "ADM-HOPE-DIR", "Edit and delete organization",
      "UAT Partner exists",
      "1. Edit name and save\n2. Delete and confirm",
      "UAT Partner",
      "Edit then delete with confirmation. Row removed."),
    r("SA-054", "Hope Events", "ADM-HOPE-EVT", "Add a Hope event",
      "On Hope Events",
      "1. Click Add Event\n2. Fill title, audience (Youth/Parents/Community), dates, optional cover\n3. Set Active\n4. Save",
      "Title: UAT Hope Event",
      "Event appears in the list with audience and dates. Toast confirms."),
    r("SA-055", "Hope Events", "ADM-HOPE-EVT", "Filter events by audience",
      "Events exist",
      "1. Click Youth, Parents, Community, then All audiences",
      "N/A",
      "List filters without reload."),
    r("SA-056", "Hope Events", "ADM-HOPE-EVT", "Edit and delete event",
      "UAT Hope Event exists",
      "1. Edit title and save\n2. Delete and confirm",
      "UAT Hope Event",
      "Edit toast. Delete confirmation then row gone."),
    r("SA-057", "Notifications", "ADM-NOTIF", "Send an in-app notification to all users",
      "On Notifications",
      "1. Click Send Notification\n2. Enter Title and Message body\n3. Audience = All Users\n4. Optional related post\n5. Send",
      "Title: UAT Notice; short message",
      "Campaign appears in history. Toast confirms. Matching app users can see it in Notifications inbox."),
    r("SA-058", "Notifications", "ADM-NOTIF", "Blank title/body blocked",
      "Send Notification dialog",
      "1. Leave Title or Message body blank\n2. Try to send",
      "Blank required fields",
      "Send is blocked with a validation message. No campaign is created."),
    r("SA-059", "Notifications", "ADM-NOTIF", "Filter, view and delete a campaign",
      "UAT Notice exists",
      "1. Filter All users / Android / iOS\n2. View the campaign\n3. Delete and confirm",
      "UAT Notice",
      "Filters work. View shows title/body/audience. Delete requires confirm then row is gone."),
    r("SA-060", "Analytics", "ADM-ANALYTICS", "Analytics command center loads",
      "Signed in as Super Admin",
      "1. Open Reach → Analytics\n2. Wait until Last updated is shown\n3. Switch Today / 7 days / 30 days",
      "N/A",
      "Cards (events, views, unique users, engagement, etc.), charts and recent activity load. Range changes refresh data. No crash."),
    r("SA-061", "Analytics", "ADM-ANALYTICS", "Refresh and export CSV",
      "On Analytics",
      "1. Click Refresh\n2. Click Export CSV\n3. Choose Last 7 days\n4. Download CSV",
      "Last 7 days",
      "Refresh updates the timestamp. CSV downloads. Toast or download proceeds without an error modal."),
    r("SA-062", "Users", "ADM-USERS", "Add an Editor with a temporary password",
      "On Users",
      "1. Click Add User\n2. Full Name, Email Address, Role = Editor\n3. Password and Confirm meeting the policy\n4. Create User",
      "Unique email; password e.g. Temp#Editor26",
      "User appears in the list as Editor. Toast confirms. That person must change password on first sign-in."),
    r("SA-063", "Users", "ADM-USERS", "Add User rejects a weak password",
      "Add User dialog",
      "1. Fill name/email/role\n2. Password = password\n3. Try Create User",
      "password",
      "Checklist shows failures. Create is blocked. No user is created."),
    r("SA-064", "Users", "ADM-USERS", "Change a user's role",
      "A non-self test user exists",
      "1. Click Change Role\n2. Set Publisher\n3. Confirm/save",
      "UAT Editor account",
      "Role updates. That account, on next sign-in, sees Publisher menus only."),
    r("SA-065", "Users", "ADM-USERS", "Cannot delete the signed-in account",
      "On Users, own row",
      "1. Find the Current account row\n2. Confirm Delete is not offered for yourself",
      "Own Super Admin row",
      "Own row shows Current account. No delete action on that row."),
    r("SA-066", "Users", "ADM-USERS", "Delete a test user",
      "Disposable UAT user exists",
      "1. Click Delete User\n2. Confirm the warning",
      "UAT user",
      "Confirmation dialog. After confirm, account is gone. Toast confirms."),
    r("SA-067", "Settings", "ADM-SETTINGS", "Edit Privacy Policy Body (English and Tagalog)",
      "On Settings",
      "1. Click Privacy Policy\n2. Edit Title/Subtitle if needed\n3. Edit Body (English) and Body (Tagalog) in the rich-text editor\n4. Add a numbered list\n5. Save",
      "Short UAT sentence in EN and TL",
      "Editor is CKEditor Classic (not a plain textarea). Numbered list shows 1, 2, 3. Toast confirms. Mobile Settings can show the update."),
    r("SA-068", "Settings", "ADM-SETTINGS", "Edit Terms of Use and Privacy Notice",
      "On Settings",
      "1. Open Terms of Use, save a minor Body change\n2. Open Privacy Notice, save a minor Body change",
      "Minor wording",
      "Both dialogs open with EN/TL Body editors. Saves succeed with toasts."),
    r("SA-069", "Settings", "ADM-SETTINGS", "Add a Kid Listo quote",
      "On Settings, Kid Listo Quotes",
      "1. Click Add Quote\n2. Enter EN and TL quote, optional attribution\n3. Save",
      "UAT quote EN/TL",
      "Quote appears in the quotes table. Toast confirms."),
    r("SA-070", "Lookups", "ADM-LOOKUPS", "Add a sky color option with subtitle and icon",
      "On Lookups; catalog What's your sky today",
      "1. Select What's your sky today\n2. Add option\n3. Labels EN/TL, subtitle EN/TL, pick an Icon, accent color\n4. Save",
      "Label: UAT Sky",
      "Option appears. Icon/subtitle are stored. Mobile journal can show the icon and subtitle under the color name."),
    r("SA-071", "Lookups", "ADM-LOOKUPS", "Edit a feeling accent color",
      "Catalog What are you feeling",
      "1. Edit a feeling\n2. Change Accent color\n3. Save",
      "Existing feeling",
      "Save succeeds. Selected chips in the app use the accent; unselected chips stay default."),
    r("SA-072", "Lookups", "ADM-LOOKUPS", "Filter catalogs and delete UAT option",
      "UAT Sky exists",
      "1. Switch catalogs on the left\n2. Delete UAT Sky and confirm",
      "UAT Sky",
      "Catalog filter works. Delete requires confirm then option is gone."),
    r("SA-073", "App Translations", "ADM-I18N", "Edit an existing English/Filipino label",
      "On App Translations",
      "1. Select Bottom menu (or another group)\n2. Search a key\n3. Edit English and Filipino\n4. Save",
      "Existing key; temporary wording then restore",
      "Save succeeds. Key is not editable. App can pick up the change without a new build (launch / ~20s / resume). Restore original wording in Remarks."),
    r("SA-074", "App Translations", "ADM-I18N", "Add Translation requires a key",
      "On App Translations",
      "1. Click Add Translation\n2. Try to save with blank key/EN/TL\n3. Cancel without creating junk data",
      "Blank fields",
      "Save is blocked. No empty key is created. This catalog is for UI chrome, not post body."),
    r("SA-075", "Word Blocking", "ADM-WORDS", "Add a blocked word",
      "On Word Blocking",
      "1. Click Add word\n2. Enter a disposable test word (min 2 characters)\n3. Active\n4. Save",
      "uatblockword",
      "Word appears as Active. Toast confirms. App replaces the whole word with **** in user text after sync."),
    r("SA-076", "Word Blocking", "ADM-WORDS", "Deactivate and delete blocked word",
      "uatblockword exists",
      "1. Edit → Inactive → save\n2. Delete and confirm",
      "uatblockword",
      "Inactive remains in the list but is not applied. Delete removes it after confirm."),
    r("SA-077", "Audit Trail", "ADM-AUDIT", "Open Audit Trail (Super Admin only)",
      "Signed in as Super Admin",
      "1. Open Administration → Audit Trail (or /admin/audit-trail)\n2. Confirm the page title and layout",
      "N/A",
      "Audit Trail loads. Metric strip (Today, This week, Unique actors, Mutations), left Category/Action filters, date range, Export CSV and the event table are visible. No Edit/Delete actions on audit rows."),
    r("SA-078", "Audit Trail", "ADM-AUDIT", "Filter by category, action, search and dates",
      "On Audit Trail; at least a few events exist",
      "1. Select Category → Users\n2. Select Action → Security\n3. Enter a known actor name in search\n4. Set From/To dates then Clear dates",
      "Known actor from an earlier Users/Auth action",
      "Each filter refreshes the table and metrics. Users shows user-related events. Security includes login, logout, password, suspend/role. Search and dates narrow rows. Clear dates restores the broader range."),
    r("SA-079", "Audit Trail", "ADM-AUDIT", "View audit event details",
      "On Audit Trail with at least one row",
      "1. Click View details (or right-click → View details)\n2. Read actor, category/action pills, IP, subject, optional before/after and request summary\n3. Close",
      "Any recent event",
      "Detail drawer opens with description, actor, IP, timestamp and sanitized properties. Passwords/tokens are not shown in clear text (redacted or omitted). Close dismisses the drawer. Row is unchanged."),
    r("SA-080", "Audit Trail", "ADM-AUDIT", "Export audit CSV",
      "On Audit Trail; optional filters applied",
      "1. Click Export CSV\n2. Open the downloaded file",
      "Current filters",
      "CSV downloads (filename includes audit-trail). Columns include created_at, actor, category, action, description, subject, IP. Content matches the filtered set (up to export limit)."),
    r("SA-081", "Audit Trail", "ADM-AUDIT", "Admin mutation is logged (Users)",
      "Super Admin; a disposable test user exists (or create one)",
      "1. On Users, change the test user's role or Suspend then Enable\n2. Open Audit Trail\n3. Filter Category → Users (or Security for suspend)\n4. Find the new event",
      "Disposable test user from SA-062+",
      "A new audit row appears with the Super Admin as actor, clear description (role change / suspended / enabled), and subject email. Viewing details does not expose passwords."),
    r("SA-082", "Audit Trail", "ADM-AUDIT", "Auth login and logout are logged",
      "Super Admin session; Audit Trail available after re-login",
      "1. Sign Out\n2. Sign in again with Super Admin\n3. Open Audit Trail → Category Auth (or Action Security)\n4. Confirm login and logout events",
      "Super Admin credentials",
      "Signed out and Signed in (or equivalent) events appear with actor email. Failed sign-in (optional: wrong password once) appears as a failed login without granting access."),
    r("SA-083", "Audit Trail", "ADM-AUDIT", "Audit Trail has no purge or edit of history",
      "On Audit Trail",
      "1. Scan Actions column and right-click menu on a row\n2. Confirm there is no Edit, Delete, Clear or Purge",
      "N/A",
      "Only View details is available. History cannot be edited or deleted from the UI."),
    r("SA-084", "Toasts / confirm", "ADM-UX", "Success uses toast; delete uses confirmation",
      "Any module with Delete (e.g. Word Blocking or Posts)",
      "1. Perform a successful save\n2. Start a Delete and click Cancel\n3. Delete and Confirm",
      "N/A",
      "Success/error appear as corner toasts, not a second success/error modal. Delete always has Cancel + confirm. Cancel leaves the record."),
]

PB = [
    r("PB-001", "Sign in", "ADM-ACCESS", "Publisher signs in",
      "Publisher account exists",
      "1. Open /admin/login\n2. Sign in as Publisher",
      "Publisher email/password",
      "Dashboard (or first allowed page) loads. Signed in as Publisher."),
    r("PB-002", "Access", "ADM-ACCESS", "Publisher sidebar shows allowed menus only",
      "Signed in as Publisher",
      "1. Scan the left menu\n2. Confirm Users, Analytics, Journal and Audit Trail are absent",
      "N/A",
      "Visible: Posts, Contests, IEC, Rehab Centers, Care Toolkit, Care Checks, Care Support, Hope Directory, Hope Events, Notifications, Settings, Lookups, App Translations, Word Blocking. Not visible: Users, Analytics, Journal, Audit Trail."),
    r("PB-003", "Access", "ADM-ACCESS", "Publisher is blocked from Users / Analytics / Journal / Audit Trail URLs",
      "Signed in as Publisher",
      "1. Open /admin/users\n2. Open /admin/analytics\n3. Open /admin/diary\n4. Open /admin/audit-trail",
      "Direct URLs",
      "Each page is refused (redirect or permission message). No user list, analytics data, journal entries, or audit events are shown."),
    r("PB-004", "Posts", "ADM-POSTS", "Publisher cannot create a post",
      "On Posts as Publisher",
      "1. Look for Create Post / Create Draft\n2. Confirm it is hidden",
      "N/A",
      "Create Draft / Create Post is not available to Publisher."),
    r("PB-005", "Posts", "ADM-POSTS", "Publisher publishes a pending post now",
      "A post is Pending review; signed in as Publisher",
      "1. Open More actions on the pending post\n2. Click Publish now",
      "Pending post",
      "Publish now is available. Status becomes Published. Toast confirms. Post can appear in the app."),
    r("PB-006", "Posts", "ADM-POSTS", "Publisher schedules and rejects",
      "Another pending post (or resubmit a draft via Editor)",
      "1. Approve & Schedule with a future time on one post\n2. Reject another pending post with review notes",
      "Future time; review notes",
      "Schedule sets Scheduled. Reject returns the post toward draft with notes. Editor can see it as not published."),
    r("PB-007", "Posts", "ADM-POSTS", "Publisher cannot archive or delete posts",
      "A published post",
      "1. Open Actions / right-click\n2. Look for Archive and Delete",
      "N/A",
      "Archive and Delete are not available to Publisher."),
    r("PB-008", "Directories", "ADM-REHAB", "Publisher can add a rehab center",
      "On Rehab Centers as Publisher",
      "1. Add Rehab Center with name, PH address suggestion and map pin\n2. Create",
      "UAT Publisher Rehab",
      "Create succeeds. Listing appears. Publisher can edit it."),
    r("PB-009", "Care / Hope / Reach", "ADM-SUPPORT", "Publisher can maintain Care Support, Hope, Notifications, Settings",
      "Publisher session",
      "1. Open Care Support and confirm Add Resource\n2. Open Hope Directory Add Organization\n3. Open Notifications Send Notification (do not need to send if policy forbids; Cancel is enough to prove the dialog)\n4. Open Settings legal tile",
      "N/A",
      "Add/edit dialogs open for Care Support, Hope Directory, Notifications and Settings. No permission error."),
    r("PB-010", "Lookups / i18n / words", "ADM-I18N", "Publisher can open Lookups, App Translations and Word Blocking",
      "Publisher session",
      "1. Open Lookups\n2. Open App Translations\n3. Open Word Blocking",
      "N/A",
      "All three pages load with their Add buttons visible."),
    r("PB-011", "Sign out", "ADM-AUTH", "Publisher signs out",
      "Signed in as Publisher",
      "1. Account menu → Sign Out",
      "N/A",
      "Returns to sign-in. Session ended."),
]

ED = [
    r("ED-001", "Sign in", "ADM-ACCESS", "Editor signs in",
      "Editor account exists",
      "1. Open /admin/login\n2. Sign in as Editor",
      "Editor email/password",
      "Admin loads. Signed in as Editor."),
    r("ED-002", "Access", "ADM-ACCESS", "Editor sidebar is Posts only",
      "Signed in as Editor",
      "1. Scan the left menu",
      "N/A",
      "Content → Posts is available. Contests, IEC, Rehab, Care, Hope, Reach, Users, Settings, Lookups, Translations, Word Blocking, Analytics, Journal, Audit Trail are not in the menu."),
    r("ED-003", "Access", "ADM-ACCESS", "Editor is blocked from other admin URLs",
      "Editor session",
      "1. Open /admin/users\n2. Open /admin/contests\n3. Open /admin/settings\n4. Open /admin/audit-trail",
      "Direct URLs",
      "Access denied or redirected. No other module data (including audit events) is shown."),
    r("ED-004", "Posts", "ADM-POSTS", "Editor creates a draft (not published)",
      "On Posts as Editor",
      "1. Click Create Draft\n2. Fill Title, Body, Category\n3. Create Draft",
      "Title: UAT Editor Draft",
      "Button/dialog is Create Draft (not Create & Publish). Post is saved as Draft and is NOT visible in the mobile app."),
    r("ED-005", "Posts", "ADM-POSTS", "Editor submits draft for review",
      "UAT Editor Draft exists",
      "1. More actions → Submit",
      "Draft post",
      "Status becomes Pending review. Editor cannot Publish now."),
    r("ED-006", "Posts", "ADM-POSTS", "Editor can edit a draft only",
      "A draft and a published post exist",
      "1. Confirm Edit is available on the draft\n2. Confirm Edit is not available on a published post",
      "Draft vs published",
      "Editor can edit drafts. Published posts are view-only for Editor."),
    r("ED-007", "Posts", "ADM-POSTS", "Editor cannot publish, schedule, archive or delete",
      "Pending or published post",
      "1. Open Actions / right-click on pending and published rows",
      "N/A",
      "Publish now, Approve & Schedule, Archive and Delete are not available to Editor."),
    r("ED-008", "Posts", "ADM-POSTS", "Editor Review Insights is hidden",
      "On Posts as Editor",
      "1. Look for Review Insights",
      "N/A",
      "Review Insights (analytics) is not shown to Editor."),
    r("ED-009", "Sign out", "ADM-AUTH", "Editor signs out",
      "Editor session",
      "1. Account menu → Sign Out",
      "N/A",
      "Returns to sign-in."),
]

AV = [
    r("AV-001", "Sign in", "ADM-ACCESS", "Analytics Viewer signs in",
      "Analytics Viewer account exists",
      "1. Open /admin/login\n2. Sign in as Analytics Viewer",
      "Analytics Viewer email/password",
      "Admin loads. Signed in as Analytics Viewer."),
    r("AV-002", "Access", "ADM-ACCESS", "Analytics Viewer sidebar is Analytics only",
      "Analytics Viewer session",
      "1. Scan the left menu",
      "N/A",
      "Reach → Analytics is available. Posts, Users, Settings and other modules are not in the menu."),
    r("AV-003", "Access", "ADM-ACCESS", "Analytics Viewer is blocked from Posts, Users and Audit Trail",
      "Analytics Viewer session",
      "1. Open /admin/posts\n2. Open /admin/users\n3. Open /admin/audit-trail",
      "Direct URLs",
      "Access denied. No posts, user list, or audit events are shown."),
    r("AV-004", "Analytics", "ADM-ANALYTICS", "Charts and cards load",
      "On Analytics",
      "1. Wait for load\n2. Switch 7 days and 30 days",
      "N/A",
      "Command center cards and charts populate. Range switch refreshes data."),
    r("AV-005", "Analytics", "ADM-ANALYTICS", "Export CSV as Analytics Viewer",
      "On Analytics",
      "1. Click Export CSV\n2. Download Last 30 days",
      "Last 30 days",
      "CSV downloads successfully."),
    r("AV-006", "Sign out", "ADM-AUTH", "Analytics Viewer signs out",
      "Analytics Viewer session",
      "1. Account menu → Sign Out",
      "N/A",
      "Returns to sign-in."),
]

SCREENS = [
    ("Overview", "Sign in — /admin/login", "Welcome back, Forgot password"),
    ("Overview", "First-time register — /admin/register", "Only if no Super Admin exists"),
    ("Overview", "Dashboard — /admin/dashboard", "Quick Actions, counts, activity"),
    ("Overview", "Edit Profile — /admin/profile", "Photo, name, change password"),
    ("Content", "Posts list — /admin/posts", "Create Post/Draft, table, Actions"),
    ("Content", "Create / Edit Post dialog", "Image, Title, Body CKEditor, Category, Comments"),
    ("Content", "View Post dialog", "Body HTML, Comments (delete), Ratings, Reactions"),
    ("Content", "Contests list — /admin/contests", "Type/status filters, Add Contest"),
    ("Content", "Add / Edit Contest dialog", "Category, dates, rules"),
    ("Content", "Contest submissions dialog", "Review entries"),
    ("Content", "IEC Materials — /admin/iec-materials", "Add, preview, edit, delete"),
    ("Services", "Rehab Centers — /admin/rehab-centers", "Map, Add Rehab Center"),
    ("Services", "Add / Edit Rehab Center dialog", "Address autocomplete, map pin"),
    ("Care", "Care Toolkit — /admin/care-toolkit", "Toolkit filters, Add Question"),
    ("Care", "Add / Edit Toolkit Question dialog", "EN/TL, answer type"),
    ("Care", "Care Checks — /admin/care-checks", "Range, filters, Export CSV"),
    ("Care", "Care Support — /admin/care-support", "Category filters, Add Resource"),
    ("Care", "Add / Edit Care Support dialog", "PH phone / short codes"),
    ("Care", "Journal — /admin/diary", "Users/Skies filters, view, delete"),
    ("Hope", "Hope Directory — /admin/hope-directory", "Category filters, Add Organization"),
    ("Hope", "Add / Edit Organization dialog", "Address autocomplete, PH phone"),
    ("Hope", "Hope Events — /admin/hope-events", "Audience filters, Add Event"),
    ("Hope", "Add / Edit Event dialog", "Dates, audience, cover"),
    ("Reach", "Notifications — /admin/notifications", "Audience filters, Send Notification"),
    ("Reach", "Send Notification dialog", "Title, message, audience, related post"),
    ("Reach", "Analytics — /admin/analytics", "Range, refresh, export, charts"),
    ("Administration", "Users — /admin/users", "Metrics, filters, Add User, role, suspend, delete"),
    ("Administration", "Add / Edit User dialog", "Password policy checklist"),
    ("Administration", "User account details drawer", "View, role, force password, suspend"),
    ("Administration", "Settings — /admin/settings", "Legal tiles, Kid Listo quotes"),
    ("Administration", "Legal page Body editor", "CKEditor EN/TL"),
    ("Administration", "Lookups — /admin/lookups", "Catalogs, Add option"),
    ("Administration", "App Translations — /admin/app-translations", "Groups, edit EN/TL"),
    ("Administration", "Word Blocking — /admin/word-blocking", "Add word, active/inactive"),
    ("Administration", "Audit Trail — /admin/audit-trail", "Metrics, filters, table/cards, Export CSV"),
    ("Administration", "Audit event detail drawer", "Actor, IP, sanitized properties, before/after"),
    ("Shared", "Mega search panel", "⌘K / search pages, actions, records"),
    ("Shared", "Right-click row Actions menu", "Same commands as Actions column"),
    ("Shared", "Delete confirmation (Swal)", "Cancel + confirm"),
    ("Shared", "Success/error toast", "Corner toast, not a success modal"),
    ("Shared", "Mobile ~375 px admin shell", "☰ menu, card lists, tap targets"),
]


def build_instructions(wb):
    ws = wb.active
    ws.title = "Instructions"
    ws.sheet_properties.tabColor = GOLD
    style_info_sheet(ws)
    ws["A1"] = "DAPE-MA – Admin Control Panel User Acceptance Test Scripts"
    ws["A1"].font = title_font
    ws.merge_cells("A1:B1")
    write_kv(ws, 3, "Project", "DAPE-MA")
    write_kv(ws, 4, "Scope", "Admin Control Panel (web) — every item in the live left sidebar")
    write_kv(ws, 5, "Document status", "Version 1.1 – for UAT execution against the admin portal (includes Audit Trail)")
    write_kv(ws, 6, "Prepared by", "")
    write_kv(
        ws,
        8,
        "Modules covered",
        "Sign in / profile · Dashboard · Posts · Contests · IEC Materials · Rehab Centers · "
        "Care Toolkit · Care Checks · Care Support · Journal · Hope Directory · Hope Events · "
        "Notifications · Analytics · Users · Settings · Lookups · App Translations · Word Blocking · "
        "Audit Trail · role access",
    )
    write_kv(
        ws,
        9,
        "Out of scope",
        "Trainings (/admin/trainings) is not in the current sidebar. Flutter end-user journeys are in the Phase 1 mobile scripts. "
        "Installer / database passwords, Code Snippets and plugin updates are not this system.",
    )
    write_kv(
        ws,
        11,
        "How to use",
        "1. Each role has its own sheet: Super Admin, Publisher, Editor, Analytics Viewer. Assign testers who match that role.\n"
        "2. Execute scripts in order within a module; some depend on earlier steps (e.g. create a draft before submit).\n"
        "3. Record what actually happened in Actual Result, then choose Pass / Fail / Blocked / Not Tested from the Status dropdown.\n"
        "4. For every Fail, set a Severity (Critical / High / Medium / Low) and log the defect quoting the Script ID.\n"
        "5. The Summary sheet updates automatically from the Status column of each script sheet.\n"
        "6. Walk every screen on the Screen Checklist (desktop and ~375 px) before signing off.",
    )
    write_kv(
        ws,
        12,
        "Severity guide",
        "Critical – crash, data loss, or core function unusable. High – major function wrong, no workaround. "
        "Medium – wrong but has a workaround. Low – cosmetic or wording.",
    )
    write_kv(
        ws,
        13,
        "Acceptance guidance",
        "Recommended for acceptance when all scripts are executed, no Critical/High failures remain open, "
        "and remaining Medium/Low items are on an agreed punch list.",
    )
    write_kv(
        ws,
        15,
        "Environment / URL",
        "Admin: /admin/login on the agreed UAT host. Use provisioned role accounts from the Test Users sheet. "
        "New passwords must include at least 8 characters, upper and lower case, a number, and a symbol. "
        "Demo seeded passwords may still work for login only (complexity is not applied to the current-password field).",
    )
    write_kv(
        ws,
        16,
        "Related document",
        "DAPE-MA Administrator Manual (docs/DAPE-MA-ADMINISTRATOR-MANUAL.docx) — staff how-to for the same menus.",
    )


def build_test_users(wb):
    ws = wb.create_sheet("Test Users")
    ws.sheet_properties.tabColor = BLUE
    headers = ["User Type", "Sheet", "Account / Access", "Who should test", "Notes"]
    widths = [22, 22, 42, 36, 62]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, i, h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = wrap_center
        cell.border = thin
        ws.column_dimensions[get_column_letter(i)].width = w
    rows = [
        ("Super Admin", "Super Admin", "Provisioned Super Admin (full sidebar)",
         "Project Owner / MISD / records administrator",
         "Runs the full admin suite including Users, Analytics, Journal and Audit Trail. Create disposable UAT records then delete them."),
        ("Publisher", "Publisher", "Provisioned Publisher account",
         "Content / program staff",
         "Content + care + hope + settings. Cannot open Users, Analytics, Journal or Audit Trail. Cannot create or archive/delete posts."),
        ("Editor", "Editor", "Provisioned Editor account",
         "Content staff who only draft posts",
         "Posts only. Creates drafts and submits for review. Cannot publish, archive or delete."),
        ("Analytics Viewer", "Analytics Viewer", "Provisioned Analytics Viewer account",
         "M&E / analytics staff",
         "Analytics only. Export CSV allowed. No content editing."),
        ("App User (negative)", "Super Admin SA-011", "Mobile app account — not an admin role",
         "Same tester as Super Admin",
         "Must be denied at /admin/login. Do not use this account on admin sheets except SA-011."),
    ]
    for r_idx, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(r_idx, c, val)
            cell.font = body_font
            cell.alignment = wrap_top
            cell.border = thin
        ws.row_dimensions[r_idx].height = 48
    ws.freeze_panes = "A2"
    note = ws.cell(8, 1, "Passwords for new/reset accounts must meet policy. Do not record live production passwords in this file.")
    note.font = small_font
    ws.merge_cells("A8:E8")


def build_summary(wb, counts):
    ws = wb.create_sheet("Summary", 1)
    ws.sheet_properties.tabColor = GOLD
    ws["A1"] = "Admin UAT – Results Summary"
    ws["A1"].font = title_font
    ws.merge_cells("A1:G1")
    headers = ["User Type / Sheet", "Total Scripts", "Pass", "Fail", "Blocked", "Not Tested", "Pass Rate"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(3, i, h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin
    widths = [28, 16, 12, 12, 12, 14, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    sheets = [
        (4, "Super Admin", "Super Admin", counts["Super Admin"]),
        (5, "Publisher", "Publisher", counts["Publisher"]),
        (6, "Editor", "Editor", counts["Editor"]),
        (7, "Analytics Viewer", "Analytics Viewer", counts["Analytics Viewer"]),
    ]
    for row, label, sheet, n in sheets:
        last = n + 1
        ws.cell(row, 1, label).font = body_font
        ws.cell(row, 2, f"=COUNTA('{sheet}'!A2:A{last})")
        ws.cell(row, 3, f"=COUNTIF('{sheet}'!J2:J{last},\"Pass\")")
        ws.cell(row, 4, f"=COUNTIF('{sheet}'!J2:J{last},\"Fail\")")
        ws.cell(row, 5, f"=COUNTIF('{sheet}'!J2:J{last},\"Blocked\")")
        ws.cell(row, 6, f"=B{row}-C{row}-D{row}-E{row}")
        ws.cell(row, 7, f"=IF(B{row}=0,0,C{row}/B{row})")
        ws.cell(row, 7).number_format = "0%"
        for c in range(1, 8):
            ws.cell(row, c).border = thin
            ws.cell(row, c).font = body_font

    ws.cell(8, 1, "TOTAL").font = Font(name="Arial", size=11, bold=True, color=WHITE)
    ws.cell(8, 2, "=SUM(B4:B7)")
    ws.cell(8, 3, "=SUM(C4:C7)")
    ws.cell(8, 4, "=SUM(D4:D7)")
    ws.cell(8, 5, "=SUM(E4:E7)")
    ws.cell(8, 6, "=SUM(F4:F7)")
    ws.cell(8, 7, "=IF(B8=0,0,C8/B8)")
    ws.cell(8, 7).number_format = "0%"
    for c in range(1, 8):
        ws.cell(8, c).fill = navy_fill
        ws.cell(8, c).font = Font(name="Arial", size=11, bold=True, color=WHITE)
        ws.cell(8, c).border = thin

    ws.cell(10, 1, "Exit criteria check").font = label_font
    checks = [
        (11, "All scripts executed (none Not Tested)", '=IF(F8=0,"MET","NOT MET")'),
        (12, "No failed scripts outstanding", '=IF(D8=0,"MET","NOT MET")'),
        (13, "No blocked scripts outstanding", '=IF(E8=0,"MET","NOT MET")'),
        (14, "Screen Checklist has no blank Matches design? rows", "Complete manually"),
    ]
    for row, label, formula in checks:
        ws.cell(row, 1, label).font = body_font
        ws.cell(row, 2, formula).font = body_font
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)

    ws.cell(16, 1, "Note: Critical/High failures must be zero for acceptance; Medium/Low failures may go on an agreed punch list. Record the final decision on the UAT Sign-Off / Acceptance Form.").font = small_font
    ws.merge_cells("A16:G16")
    ws.column_dimensions["A"].width = 62


def build_screen_checklist(wb):
    ws = wb.create_sheet("Screen Checklist")
    ws.sheet_properties.tabColor = BLUE
    headers = [
        "Section",
        "Screen / dialog",
        "What to check",
        "Matches design?",
        "Controls all work?",
        "OK on ~375 px?",
        "Tester",
        "Remarks",
    ]
    widths = [18, 42, 48, 16, 18, 16, 14, 36]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, i, h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    yes_no = DataValidation(type="list", formula1='"Yes,No,N/A"', allow_blank=True, showDropDown=False)
    last = len(SCREENS) + 1
    yes_no.add(f"D2:F{last}")
    ws.add_data_validation(yes_no)
    for idx, (section, screen, check) in enumerate(SCREENS, 2):
        values = (section, screen, check, None, None, None, None, None)
        for c, val in enumerate(values, 1):
            cell = ws.cell(idx, c, val)
            cell.font = body_font
            cell.alignment = wrap_top
            cell.border = thin
            if idx % 2 == 0:
                cell.fill = zebra_fill
        ws.row_dimensions[idx].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{last}"
    note = ws.cell(last + 2, 1, "Fill Matches design? / Controls all work? / OK on ~375 px? before sign-off. Add rows if a new admin screen is added.")
    note.font = small_font
    ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=8)


def build():
    wb = Workbook()
    build_instructions(wb)
    add_script_sheet(wb, "Super Admin", SA, NAVY)
    add_script_sheet(wb, "Publisher", PB, BLUE)
    add_script_sheet(wb, "Editor", ED, "374151")
    add_script_sheet(wb, "Analytics Viewer", AV, "055498")
    build_summary(
        wb,
        {
            "Super Admin": len(SA),
            "Publisher": len(PB),
            "Editor": len(ED),
            "Analytics Viewer": len(AV),
        },
    )
    build_test_users(wb)
    build_screen_checklist(wb)

    order = [
        "Instructions",
        "Summary",
        "Test Users",
        "Super Admin",
        "Publisher",
        "Editor",
        "Analytics Viewer",
        "Screen Checklist",
    ]
    wb._sheets = [wb[name] for name in order]

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(
        "scripts",
        len(SA) + len(PB) + len(ED) + len(AV),
        "SA",
        len(SA),
        "PB",
        len(PB),
        "ED",
        len(ED),
        "AV",
        len(AV),
        "screens",
        len(SCREENS),
    )


if __name__ == "__main__":
    build()
