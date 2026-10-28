#!/usr/bin/env python3
"""Build DAPE-MA-ADMINISTRATOR-MANUAL.docx from Markdown chapters + figures."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x12, 0x3A, 0x60)
BLUE = RGBColor(0x05, 0x54, 0x98)
GOLD = RGBColor(0xFB, 0xD1, 0x16)
RED = RGBColor(0xCE, 0x20, 0x28)
BLACK = RGBColor(0x0A, 0x0A, 0x0A)
MUTED = RGBColor(0x37, 0x41, 0x51)

HERE = Path(__file__).resolve().parent
CHAPTERS_DIR = HERE / "chapters"
FIGURES_DIR = HERE / "figures"
OUT_PATH = HERE.parent / "DAPE-MA-ADMINISTRATOR-MANUAL.docx"
SEAL = FIGURES_DIR / "ddb-seal.png"

CHAPTER_FILES = [
    "00-where-to-find.md",
    "01-sign-in.md",
    "02-duties.md",
    "03-lists.md",
    "04-dashboard.md",
    "05-posts.md",
    "06-contests.md",
    "07-iec.md",
    "08-rehab.md",
    "09-care-toolkit.md",
    "10-care-checks.md",
    "11-care-support.md",
    "12-journal.md",
    "13-hope-directory.md",
    "14-hope-events.md",
    "15-notifications.md",
    "16-analytics.md",
    "17-users.md",
    "18-settings.md",
    "19-lookups.md",
    "20-translations.md",
    "21-word-blocking.md",
    "22-troubleshooting.md",
    "23-quick-reference.md",
]


def set_run_font(run, name="Calibri", size=11, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), kwargs[edge].get("val", "single"))
            element.set(qn("w:sz"), str(kwargs[edge].get("sz", 4)))
            element.set(qn("w:color"), kwargs[edge].get("color", "123A60"))
            tc_borders.append(element)
    tc_pr.append(tc_borders)


def add_page_border(section, color="123A60"):
    sect_pr = section._sectPr
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:space"), "24")
        border.set(qn("w:color"), color)
        pg_borders.append(border)
    sect_pr.append(pg_borders)


def add_horizontal_line(paragraph, color="FBD116"):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    r = run._r
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text_run = OxmlElement("w:t")
    text_run.text = "Right-click and choose Update Field to refresh this table of contents."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(text_run)
    r.append(fld_end)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, space_before, space_after in (
        ("Title", 28, NAVY, 0, 6),
        ("Heading 1", 18, NAVY, 18, 8),
        ("Heading 2", 14, BLUE, 14, 6),
        ("Heading 3", 12, NAVY, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.font.color.rgb = NAVY
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(12)


def add_cover(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    add_page_border(section)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(24)

    if SEAL.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(SEAL), width=Inches(1.35))

    agency = doc.add_paragraph()
    agency.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = agency.add_run("DANGEROUS DRUGS BOARD")
    set_run_font(run, size=12, bold=True, color=NAVY)

    gold = doc.add_paragraph()
    gold.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_horizontal_line(gold)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    run = title.add_run("DAPE-MA")
    set_run_font(run, size=32, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("ADMINISTRATOR MANUAL")
    set_run_font(run, size=22, bold=True, color=BLUE)

    version = doc.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = version.add_run("VERSION 1.0")
    set_run_font(run, size=14, bold=True, color=NAVY)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Staff how-to guide for the DAPE-MA Admin Control Panel")
    set_run_font(run, size=11, italic=True, color=MUTED)

    for _ in range(3):
        doc.add_paragraph()

    addr = doc.add_paragraph()
    addr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = addr.add_run("DANGEROUS DRUGS BOARD")
    set_run_font(run, size=11, bold=True, color=NAVY)

    for line in (
        "3rd Floor DDB-PDEA Bldg.,",
        "National Government Center,",
        "NIA Road, Brgy. Pinyahan",
        "1100 Quezon City, Philippines",
        "www.ddb.gov.ph",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        set_run_font(run, size=11, color=MUTED)

    doc.add_page_break()


def add_contents_page(doc: Document, headings: list[tuple[int, str]]):
    p = doc.add_paragraph()
    run = p.add_run("Contents")
    set_run_font(run, size=22, bold=True, color=NAVY)
    add_horizontal_line(p)

    hint = doc.add_paragraph()
    add_toc_field(hint)
    hint.paragraph_format.space_after = Pt(12)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    hdr = table.rows[0].cells
    hdr[0].text = "Chapter"
    hdr[1].text = "Page"
    for cell in hdr:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, "123A60")

    for level, text in headings:
        if level > 2:
            continue
        row = table.add_row().cells
        prefix = "    " if level == 2 else ""
        row[0].text = f"{prefix}{text}"
        row[1].text = ""
        for paragraph in row[0].paragraphs:
            for run in paragraph.runs:
                set_run_font(run, size=11, bold=(level == 1), color=NAVY if level == 1 else MUTED)

    doc.add_paragraph()
    doc.add_page_break()


def add_inline_runs(paragraph, text: str):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=11, bold=True, color=BLACK)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=10, color=BLUE)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=11, color=BLACK)


def add_callout(doc: Document, kind: str, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    fill = "F8E7E8" if kind == "Important" else "EEF5FB"
    shade_cell(cell, fill)
    border_color = "CE2028" if kind == "Important" else "055498"
    set_cell_border(
        cell,
        top={"val": "single", "sz": 12, "color": border_color},
        left={"val": "single", "sz": 24, "color": border_color},
        bottom={"val": "single", "sz": 12, "color": border_color},
        right={"val": "single", "sz": 12, "color": border_color},
    )
    paragraph = cell.paragraphs[0]
    label = paragraph.add_run(f"{kind}. ")
    set_run_font(label, size=11, bold=True, color=RED if kind == "Important" else BLUE)
    add_inline_runs(paragraph, text)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def add_image(doc: Document, rel_path: str, caption: str):
    name = Path(rel_path).name
    path = (FIGURES_DIR / name).resolve()
    if not path.exists():
        p = doc.add_paragraph()
        run = p.add_run(f"[Missing figure: {name}]")
        set_run_font(run, size=10, italic=True, color=RED)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.2))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.style = "Caption"
    run = cap.add_run(caption)
    set_run_font(run, size=10, italic=True, color=NAVY)


def add_md_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_inline_runs(paragraph, value.strip())
            if r_idx == 0:
                shade_cell(cell, "123A60")
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.bold = True
            elif r_idx % 2 == 0:
                shade_cell(cell, "F3F6FA")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def collect_headings(markdown: str) -> list[tuple[int, str]]:
    headings = []
    for line in markdown.splitlines():
        if line.startswith("# "):
            headings.append((1, line[2:].strip()))
        elif line.startswith("## "):
            headings.append((2, line[3:].strip()))
    return headings


def render_markdown(doc: Document, markdown: str):
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        image = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
        if image:
            add_image(doc, image.group(2), image.group(1))
            i += 1
            continue

        if line.startswith("> "):
            body = line[2:].strip()
            kind = "Tip"
            if body.startswith("**Important.**"):
                kind = "Important"
                body = body.replace("**Important.**", "", 1).strip()
            elif body.startswith("**Tip.**"):
                kind = "Tip"
                body = body.replace("**Tip.**", "", 1).strip()
            add_callout(doc, kind, body)
            i += 1
            continue

        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip()
                cells = [c.strip() for c in raw.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
                    rows.append(cells)
                i += 1
            add_md_table(doc, rows)
            continue

        ordered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if ordered:
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, ordered.group(2))
            i += 1
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, line[2:].strip())
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, line.strip())
        i += 1


def build():
    chapters = []
    for name in CHAPTER_FILES:
        path = CHAPTERS_DIR / name
        chapters.append(path.read_text(encoding="utf-8"))

    all_md = "\n\n".join(chapters)
    headings = collect_headings(all_md)

    doc = Document()
    configure_styles(doc)
    add_cover(doc)
    add_contents_page(doc, headings)

    for markdown in chapters:
        render_markdown(doc, markdown)

    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("DAPE-MA Administrator Manual  ·  Version 1.0  ·  Dangerous Drugs Board")
    set_run_font(run, size=8, color=NAVY)

    header = doc.sections[0].header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("DAPE-MA  ·  Admin Control Panel")
    set_run_font(run, size=9, bold=True, color=BLUE)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_PATH))
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    build()
